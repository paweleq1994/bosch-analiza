#!/usr/bin/env python3
"""Build a source-linked Excel archive from the exact validated portal snapshot."""
from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/assets/data.json"
DESTINATION = ROOT / "docs/downloads/bosch-kompendium.xlsx"
CHECKSUM = ROOT / "docs/downloads/bosch-kompendium.source.sha256"
COLUMNS = {
    "Zawodnicy": [
        "PLP ID", "Zawodnik", "M Bosch", "G Bosch", "A Bosch",
        "G+A Bosch", "MVP Bosch", "ŻK Bosch", "CK bezpośrednie Bosch",
        "M PLP kariera", "G PLP kariera", "A PLP kariera",
        "Kompletność Bosch", "Źródło profilu", "Wykluczenia 2ŻK — protokoły ligowe",
    ],
    "Sezony": [
        "ID", "Sezon", "M", "W", "R", "P", "Punkty",
        "Gole zdobyte", "Gole stracone", "Bilans bramek",
        "Miejsce", "Źródło",
    ],
    "Mecze": [
        "ID meczu", "Sezon", "Data", "Rywal", "ID rywala",
        "Gole Bosch", "Gole rywala", "Bilans", "Status protokołu", "Protokół PLP",
    ],
    "Występy": [
        "ID meczu", "Data", "ID zawodnika", "Zawodnik",
        "Występ", "Gole", "Asysty", "MVP", "ŻK", "CK bezpośrednie",
        "Wykluczenie 2ŻK", "Wykluczenia łącznie w meczu", "Protokół PLP",
    ],
    "Bramki": [
        "ID meczu", "Sezon", "Data", "Zawodnik", "ID zawodnika",
        "Gole", "Minuty", "Niepewna minuta", "Protokół PLP",
    ],
    "Braki danych": [
        "ID zawodnika", "Zawodnik", "Brak pełnej kariery",
        "Brak zdjęcia", "Niepełny bilans sezonów",
        "Źródło tożsamości",
    ],
}
HEADER_FILL = PatternFill("solid", fgColor="172D49")
SECTION_FILL = PatternFill("solid", fgColor="E7F1F5")
HEADER_FONT = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
SOURCE_FONT = Font(name="Aptos", size=10, color="087A56", underline="single")
VALUE_FONT = Font(name="Aptos", size=10, color="166534")
STATIC_FONT = Font(name="Aptos", size=10, color="334155")
REVIEW_FONT = Font(name="Aptos", size=10, color="B45309")
FORMULA_FONT = Font(name="Aptos", size=10, color="111827")


def cell_value(value):
    if isinstance(value, (bool, int, float)) or value is None:
        return value

    text = str(value)

    if text.startswith(("=", "+", "-", "@")):
        return "'" + text

    return text


def add_rows(workbook: Workbook, name: str, rows: list[list], source_column: int | None = None) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(COLUMNS[name])
    sheet.freeze_panes = "C2"
    sheet.sheet_view.showGridLines = False

    for index, head in enumerate(COLUMNS[name], start=1):
        cell = sheet.cell(1, index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    sheet.row_dimensions[1].height = 34

    for row in rows:
        sheet.append([cell_value(value) for value in row])

    last = len(COLUMNS[name])
    sheet.auto_filter.ref = f"A1:{get_column_letter(last)}{sheet.max_row}"

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = VALUE_FONT if isinstance(cell.value, (int, float)) else STATIC_FONT
            cell.alignment = Alignment(vertical="center")

        if source_column is not None:
            item = row[source_column - 1]
            if isinstance(item.value, str) and item.value.startswith("https://"):
                item.hyperlink = item.value
                item.font = SOURCE_FONT

    for index, heading in enumerate(COLUMNS[name], start=1):
        width = 36 if "Źródło" in heading or "Protokół" in heading else 26 if heading in {
            "Zawodnik", "Rywal", "Sezon", "Źródło tożsamości"
        } else 17
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.outlinePr.summaryBelow = True


def career_source(player: dict) -> str:
    career = player.get("career_totals") or {}
    return (
        str(career.get("source_url") or "")
        or str((career.get("source_urls") or [""])[0])
        or str(player.get("profile_url") or "")
        or str(((player.get("verified_identity_source") or {}).get("evidence_urls") or [""])[0])
    )


def build_workbook(database: dict) -> Workbook:
    workbook = Workbook()
    workbook.properties.identifier = exporter_signature()
    info = workbook.active
    info.title = "Informacje"
    info.sheet_view.showGridLines = False
    info.column_dimensions["A"].width = 33
    info.column_dimensions["B"].width = 89
    info.append(["Bosch Service Tech-Car — oficjalne archiwum PLP", ""])
    info["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="183B56")
    info.merge_cells("A1:B1")
    info.row_dimensions[1].height = 38
    notes = [
        ("Źródło", "https://podlaskaliga.pl/"),
        ("Kompendium", "https://paweleq1994.github.io/bosch-analiza/"),
        ("Repozytorium", "https://github.com/paweleq1994/bosch-analiza"),
        ("Wersja danych", str((database.get("meta") or {}).get("source_updated") or "brak daty")),
        ("Zawodników", len(database["players"])),
        ("Meczów", len(database["matches"])),
        ("Uwaga", "PLP kariera = wszystkie kluby. Bosch = tylko mecze w naszej drużynie."),
        ("Uwaga", "Puste pola oznaczają brak potwierdzenia, NIE zero."),
        ("Uwaga", "Statystyki meczowe pokazują tylko rozpoznane protokoły."),
        ("Uwaga", "Bramki w arkuszu Bramki pochodzą z rozliczonych protokołów strzelców."),
        ("Uwaga", "Formuły G+A i bilansu są obliczane w Excelu przy otwarciu."),
    ]

    for key, value in notes:
        info.append([key, cell_value(value)])

    for player in database["players"]:
        review = player.get("public_protocol_review")
        if not review:
            continue
        totals = review["totals"]
        career = player.get("career_totals") or {}
        reconstructed = career.get("reconstructed")
        prefix = "" if reconstructed else "≥"
        info.append([("Rekonstrukcja PLP — " if reconstructed else "Minimum PLP — ") + player["name"],
                     ", ".join(f"{field}: {prefix}{totals[field]}" for field in ("apps", "goals", "assists", "motm", "yellow", "red"))
                     + ". " + (career["note"] if reconstructed else "Pełna kariera niepotwierdzona; odczyt publicznych protokołów.")])
        for match in review["matches"]:
            if any(str(league["id"]) == match["event_id"] for league in database["matches"]):
                continue
            info.append([f"{player['name']} — PLP {match['event_id']}", match["source_url"]])
            info.cell(info.max_row, 2).hyperlink = match["source_url"]
            info.cell(info.max_row, 2).font = SOURCE_FONT
    for cell in info["B"][1:]:
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for cell in info["A"][1:]:
        cell.fill = SECTION_FILL
        cell.font = Font(name="Aptos", size=11, bold=True, color="17324D")

    info.freeze_panes = "A2"

    players = database["players"]
    players_by_id = {str(player["id"]): player for player in players}
    photo_ids = {
        str(item["id"])
        for item in ((database.get("media") or {}).get("subjects") or [])
        if item.get("kind") == "players" and item.get("photos")
    }
    player_rows = []
    for player in sorted(players, key=lambda item: item["name"].casefold()):
        career = player.get("career_totals") or {}
        player_rows.append([
            str(player["id"]), player["name"], player.get("apps"), player.get("goals"),
            player.get("assists"), None, player.get("motm"),
            player.get("yellow"), player.get("red"),
            career.get("apps"), career.get("goals"), career.get("assists"),
            "Rekonstrukcja z protokołów" if career.get("reconstructed") else "Potwierdzone minimum" if player.get("stats_minimum") or not player.get("bosch_totals") else "Zweryfikowane sezony",
            career_source(player),
            player.get("second_yellow"),
        ])
    add_rows(workbook, "Zawodnicy", player_rows, source_column=14)
    for row_number in range(2, workbook["Zawodnicy"].max_row + 1):
        cell = workbook["Zawodnicy"][f"F{row_number}"]
        cell.value = f'=IF(AND(ISNUMBER(D{row_number}),ISNUMBER(E{row_number})),SUM(D{row_number}:E{row_number}),"")'
        cell.font = FORMULA_FONT

    league_rows = []
    for row in database["seasons"]:
        league_rows.append([
            str(row["id"]), row["name"], row.get("matches"), row.get("wins"),
            row.get("draws"), row.get("losses"), row.get("points"),
            row.get("gf"), row.get("ga"), None, row.get("position"),
            "https://podlaskaliga.pl/",
        ])
    add_rows(workbook, "Sezony", league_rows, source_column=12)
    for index in range(2, workbook["Sezony"].max_row + 1):
        cell = workbook["Sezony"][f"J{index}"]
        cell.value = f'=IF(AND(ISNUMBER(H{index}),ISNUMBER(I{index})),H{index}-I{index},"")'
        cell.font = FORMULA_FONT

    match_rows = []
    event_rows = []
    scorer_rows = []
    for match in sorted(database["matches"], key=lambda item: (item["date"], str(item["id"]))):
        match_id = str(match["id"])
        source = str(match.get("event_link") or "")
        protocol = match.get("goal_protocol") or {}
        verified = bool(protocol.get("verified"))
        match_rows.append([
            match_id, match["season"], match["date"], match["opponent"],
            str(match.get("opponent_id") or ""), match["gf"], match["ga"],
            None, "Zweryfikowany" if verified else "Oczekuje", source,
        ])
        for item in match.get("player_performance") or []:
            pid = str(item["player_id"])
            official_name = players_by_id.get(pid, {}).get("name") or item.get("name") or f"PLP #{pid}"
            event_rows.append([
                match_id, match["date"], pid, official_name,
                item.get("appearance"), item.get("goals"), item.get("assists"),
                item.get("motm"), item.get("yellow"), item.get("red"),
                item.get("second_yellow"), item.get("dismissals"),
                match.get("player_performance_source") or source,
            ])
        for scorer in protocol.get("scorers") or []:
            scorer_rows.append([
                match_id, match["season"], match["date"], scorer["name"],
                str(scorer["player_id"]), scorer["goals"],
                ", ".join(str(v) for v in scorer.get("minutes") or []),
                "TAK" if scorer.get("minute_uncertain") else "NIE",
                protocol.get("source_url") or source,
            ])

    add_rows(workbook, "Mecze", match_rows, source_column=10)
    for index in range(2, workbook["Mecze"].max_row + 1):
        cell = workbook["Mecze"][f"H{index}"]
        cell.value = f'=IF(AND(ISNUMBER(F{index}),ISNUMBER(G{index})),F{index}-G{index},"")'
        cell.font = FORMULA_FONT
    add_rows(workbook, "Występy", event_rows, source_column=13)
    add_rows(workbook, "Bramki", scorer_rows, source_column=9)

    missing_rows = []
    for player in sorted(players, key=lambda item: item["name"].casefold()):
        missing_career = not bool(player.get("career_totals"))
        missing_photo = str(player["id"]) not in photo_ids
        unassigned = (player.get("season_coverage") or {}).get("unassigned") or {}
        incomplete = any(value is None or value > 0 for value in unassigned.values())
        if not (missing_career or missing_photo or incomplete):
            continue
        missing_rows.append([
            str(player["id"]), player["name"],
            "TAK" if missing_career else "NIE",
            "TAK" if missing_photo else "NIE",
            "TAK" if incomplete else "NIE",
            str((player.get("verified_identity_source") or {}).get("type") or "PLP"),
        ])
    add_rows(workbook, "Braki danych", missing_rows)
    for row in workbook["Braki danych"].iter_rows(min_row=2):
        for cell in row[2:5]:
            if cell.value == "TAK":
                cell.font = REVIEW_FONT

    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    return workbook


def exporter_signature() -> str:
    return sha256(Path(__file__).read_text(encoding="utf-8").encode("utf-8")).hexdigest()


def main() -> None:
    # Git normalizes text to LF; Windows checkouts may use CRLF. Both represent
    # the same committed snapshot and must share one downloadable workbook.
    raw = SOURCE.read_text(encoding="utf-8").encode("utf-8")
    digest = sha256(raw).hexdigest()
    if DESTINATION.exists() and CHECKSUM.exists() and CHECKSUM.read_text().strip() == digest:
        cached = load_workbook(DESTINATION, read_only=True)
        try:
            current_exporter = cached.properties.identifier == exporter_signature()
        finally:
            cached.close()
        if current_exporter:
            print("XLSX export unchanged", flush=True)
            return

    database = json.loads(raw)
    workbook = build_workbook(database)
    DESTINATION.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(DESTINATION)
    CHECKSUM.write_text(digest + "\n", encoding="ascii")
    verify = load_workbook(DESTINATION, read_only=True, data_only=False)
    try:
        assert verify["Zawodnicy"].max_row == len(database["players"]) + 1
        assert verify["Mecze"].max_row == len(database["matches"]) + 1
    finally:
        verify.close()
    print(
        f"XLSX exported: {len(database['players'])} players, "
        f"{len(database['matches'])} matches, sheets={len(verify.sheetnames)}",
        flush=True,
    )


if __name__ == "__main__":
    main()
